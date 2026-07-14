#!/usr/bin/env python3
"""Importa ZONA CENTRO.xlsx al sistema Hospedaje (evento 8).

Mapeo acordado:
- N° (MIRH-n)        -> folio_anterior
- NOMBRE ANFITRION   -> nombre_dueno
- CELULAR 1          -> telefono_dueno (10 dígitos); CELULAR 2 -> comentarios "Cel 2: ..."
- CALLE+N°EXT+N°INT  -> calle_numero ("Arquímedes 630 Int. A"; INT 0 se ignora)
- colonia fija "Hermosa Provincia", CP "44770", estado "Jalisco"
- TIPO CASA          -> tenencia (Propia / Renta->Rentada; raro/vacío -> sin valor)
- CUPO               -> capacidad (MIRH-130 "10?" -> 10 + nota en comentarios)
- COMENTARIOS        -> comentarios (se conserva íntegro salvo placeholders O/Oo/0/0.0/.)
                        + extracción: vulnerabilidad "Planta alta / escaleras",
                        perfil "Solo mujeres" / "Solo hombres" (58,189,439)
- GPS                -> lat/lng (coords directas o resolviendo link Google);
                        texto -> comentarios "Ubicación: ..."; link irresoluble -> "Mapa: <link>"
- STATUS PENDIENTE   -> se omiten (van a pendientes_zona_centro.csv)
"""
import openpyxl, re, sys, json, csv, unicodedata, urllib.request, urllib.parse
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

XLSX = '/Users/davidchaveznunez/Documents/General programming/Hospedaje/Inputs/ZONA CENTRO.xlsx'
PENDIENTES_CSV = '/Users/davidchaveznunez/Documents/General programming/Hospedaje/Inputs/pendientes_zona_centro.csv'
HERE = Path(__file__).parent
CACHE = HERE / 'gps_cache.json'
API = 'https://lldmhospedaje.tech/api'
EVENTO_ID = 8

JUNK = {'', 'o', 'oo', '0', '0.0', '.', 'none'}
HOMBRES_FOLIOS = {'MIRH-58', 'MIRH-189', 'MIRH-439'}

def norm(s):
    s = unicodedata.normalize('NFD', str(s or '').lower())
    return ''.join(c for c in s if unicodedata.category(c) != 'Mn')

RE_ALTA = re.compile(r'escalera|escalon|planta alta|p\.?\s*a(?![a-z])|(segund|tercer|2d?a|3e?r)a?\s*(planta|piso|nivel)|piso\s*[23]|[23]\s*planta|2do piso|altos')
RE_MUJ = re.compile(r'(solo|puras?|enviar|mandar)[^.,;]*?(mujeres|hermanas|hnas|senoritas)|mujeres solas|hermanas solas|puras mujeres')
RE_COORD = re.compile(r'^(-?\d+\.\d+)\s*,\s*(-?\d+\.\d+)')

def parse_rows():
    ws = openpyxl.load_workbook(XLSX, data_only=True)['Zona Centro']
    out = []
    for row in ws.iter_rows(min_row=2, max_row=482, max_col=13):
        v = [c.value for c in row]
        if v[0] and str(v[0]).startswith('MIRH'):
            out.append(v)
    return out

def limpia_num(v):
    s = re.sub(r'\.0$', '', str(v).strip()) if v is not None else ''
    return re.sub(r'\D', '', s)

def fmt_ext(v):
    s = str(v).strip() if v is not None else ''
    s = re.sub(r'\.0$', '', s)
    return s

def calle_numero(v):
    calle = ' '.join(str(v[5]).split())
    ext = fmt_ext(v[6])
    interior = fmt_ext(v[7])
    s = f'{calle} {ext}'.strip()
    if interior and interior not in ('0', '0.0'):
        s += f' Int. {interior.upper()}'
    return s

def gps_kind(s):
    s = (s or '').strip()
    if not s or s in ('0.0', '0', '.', 'No'): return 'vacio'
    if 'goo.gl' in s or 'google.com/maps' in s: return 'link'
    if RE_COORD.match(s): return 'coord'
    return 'texto'

def extrae_coords_url(url):
    for pat in (r'!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)', r'/search/(-?\d+\.\d+),\s*\+?(-?\d+\.\d+)',
                r'[?&]q=(-?\d+\.\d+)\s*,\s*\+?(-?\d+\.\d+)', r'@(-?\d+\.\d+),(-?\d+\.\d+)'):
        m = re.search(pat, urllib.parse.unquote(url))
        if m:
            return float(m.group(1)), float(m.group(2))
    return None

def resolver_link(url):
    import subprocess
    url = url.strip().split()[0]
    ua = 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36'
    try:
        final = subprocess.run(
            ['curl', '-s', '-o', '/dev/null', '-w', '%{url_effective}', '-L', '-m', '20', '-A', ua, url],
            capture_output=True, text=True, timeout=25).stdout.strip()
        coords = extrae_coords_url(final)
        if coords: return coords
        html = subprocess.run(['curl', '-s', '-L', '-m', '20', '-A', ua, url],
                              capture_output=True, text=True, timeout=25).stdout
        for m in re.finditer(r'\[(-?\d{1,2}\.\d{4,}),(-?\d{1,3}\.\d{4,})\]', html):
            lat, lng = float(m.group(1)), float(m.group(2))
            if 18 < lat < 23 and -106 < lng < -101:  # sanity: Jalisco y alrededores
                return lat, lng
    except Exception:
        pass
    return None

def main():
    dry = '--dry-run' in sys.argv
    rows = parse_rows()
    pend = [v for v in rows if v[12] and 'PEND' in str(v[12]).upper()]
    conf = [v for v in rows if not (v[12] and 'PEND' in str(v[12]).upper())]
    print(f'filas: {len(rows)} | confirmadas: {len(conf)} | pendientes: {len(pend)}')

    # CSV de pendientes
    with open(PENDIENTES_CSV, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['folio', 'nombre', 'celular1', 'celular2', 'calle', 'num_ext', 'num_int', 'tipo_casa', 'cupo', 'comentarios', 'gps', 'status'])
        for v in pend:
            w.writerow([str(v[0])] + [str(x) if x is not None else '' for x in (v[1], limpia_num(v[2]), limpia_num(v[3]), v[5], fmt_ext(v[6]), fmt_ext(v[7]), v[8], v[9], v[10], v[11], v[12])])
    print(f'pendientes guardadas en {PENDIENTES_CSV}')

    # resolver links GPS (con caché)
    cache = json.loads(CACHE.read_text()) if CACHE.exists() else {}
    links = [str(v[11]).strip() for v in conf if gps_kind(str(v[11] or '')) == 'link']
    faltan = [u for u in links if u not in cache]
    print(f'links GPS: {len(links)} | por resolver: {len(faltan)}')
    if faltan:
        with ThreadPoolExecutor(max_workers=10) as ex:
            for url, res in zip(faltan, ex.map(resolver_link, faltan)):
                cache[url] = res
        CACHE.write_text(json.dumps(cache))
    resueltos = sum(1 for u in links if cache.get(u))
    print(f'links resueltos a coordenadas: {resueltos}/{len(links)}')

    # armar payloads
    payloads = []
    stats = {'con_pin': 0, 'tenencia': {}, 'vuln': 0, 'muj': 0, 'hom': 0, 'sin_tel': 0}
    conf.sort(key=lambda v: int(re.sub(r'\D', '', str(v[0]))))
    for v in conf:
        folio = str(v[0]).strip()
        n = norm(v[10])
        comentario = str(v[10] or '').strip()
        partes = [] if norm(comentario).strip() in JUNK else [comentario]

        cel2 = limpia_num(v[3])
        if cel2: partes.append(f'Cel 2: {cel2}')

        cupo_raw = str(v[9]).strip() if v[9] is not None else ''
        if folio == 'MIRH-130':
            capacidad = 10
            partes.append(f'Cupo por confirmar ({cupo_raw})')
        else:
            capacidad = int(float(cupo_raw))

        gps = str(v[11] or '').strip()
        kind = gps_kind(gps)
        lat = lng = None
        if kind == 'coord':
            m = RE_COORD.match(gps); lat, lng = float(m.group(1)), float(m.group(2))
        elif kind == 'link':
            res = cache.get(gps)
            if res: lat, lng = res
            else: partes.append(f'Mapa: {gps.split()[0]}')
        elif kind == 'texto':
            partes.append(f'Ubicación: {gps}')
        if lat: stats['con_pin'] += 1

        tipo = str(v[8] or '').strip()
        tenencia = {'Propia': 'Propia', 'Renta': 'Rentada'}.get(tipo, '')
        stats['tenencia'][tenencia or 'NULL'] = stats['tenencia'].get(tenencia or 'NULL', 0) + 1

        vuln = ['Planta alta / escaleras'] if RE_ALTA.search(n) else []
        if vuln: stats['vuln'] += 1
        if folio in HOMBRES_FOLIOS:
            perfil = ['Solo hombres']; stats['hom'] += 1
        elif RE_MUJ.search(n):
            perfil = ['Solo mujeres']; stats['muj'] += 1
        else:
            perfil = []

        tel = limpia_num(v[2])
        if not tel: stats['sin_tel'] += 1

        p = {
            'evento_id': str(EVENTO_ID),
            'nombre_dueno': ' '.join(str(v[1]).split()),
            'telefono_dueno': tel,
            'calle_numero': calle_numero(v),
            'colonia': 'Hermosa Provincia',
            'codigo_postal': '44770',
            'estado': 'Jalisco',
            'referencias': '',
            'capacidad': str(capacidad),
            'tenencia': tenencia,
            'comentarios': ' · '.join(partes),
            'folio_anterior': folio,
            'servicios': '[]',
            'vulnerabilidades': json.dumps(vuln, ensure_ascii=False),
            'notas_vulnerabilidad': '',
            'perfil_sugerido': json.dumps(perfil, ensure_ascii=False),
        }
        if lat: p['lat'] = str(lat); p['lng'] = str(lng)
        payloads.append(p)

    print(f"payloads: {len(payloads)} | con pin: {stats['con_pin']} | tenencia: {stats['tenencia']}")
    print(f"vuln planta alta: {stats['vuln']} | solo mujeres: {stats['muj']} | solo hombres: {stats['hom']} | sin teléfono: {stats['sin_tel']}")
    (HERE / 'payloads_preview.json').write_text(json.dumps(payloads[:8] + payloads[-3:], ensure_ascii=False, indent=1))
    if dry:
        print('DRY RUN: no se envió nada. Muestra en payloads_preview.json')
        return

    import subprocess
    token = (HERE / 'token.txt').read_text().strip()
    errores = []
    for i, p in enumerate(payloads, 1):
        cmd = ['curl', '-s', '-o', '/dev/null', '-w', '%{http_code}', '-m', '30',
               '-X', 'POST', f'{API}/hogares', '-H', f'Authorization: Bearer {token}']
        for k, val in p.items():
            cmd += ['--form-string', f'{k}={val}']
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=40)
        if r.stdout.strip() != '201':
            errores.append((p['folio_anterior'], f'HTTP {r.stdout.strip()}', ''))
        if i % 50 == 0: print(f'  {i}/{len(payloads)}...')
    print(f'enviados: {len(payloads) - len(errores)} | errores: {len(errores)}')
    for f, e, d in errores[:20]: print('  ERROR', f, e, d)

if __name__ == '__main__':
    main()
